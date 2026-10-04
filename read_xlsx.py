from openpyxl import load_workbook

workbook = load_workbook("abc.xlsx")
sheet = workbook['Sheet1']

for row in sheet1.iter_rows(values_only= True):
    print(row)